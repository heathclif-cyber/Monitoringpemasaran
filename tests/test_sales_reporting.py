import io
import unittest
from datetime import date
from unittest.mock import patch
from fastapi import HTTPException
from openpyxl import load_workbook
import test_ba_workflow as fixture
import models, schemas
from api.r_laporan import _build_laporan_rows
from api.r_do import create_do, delete_do
from api.r_invoice import delete_invoice
from api.r_pembayaran import create_pembayaran
from services.document_flow import build_document_flow
from services.laporan_ho_export import generate_laporan_ho_xlsx, LOKAL_ROW_BY_KEY


class SalesReportingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        fixture.BAWorkflowTests.setUpClass()

    def setUp(self):
        self.f = fixture.BAWorkflowTests()
        self.f.setUp()
        self.db = self.f.db

    def tearDown(self):
        self.f.tearDown()

    def rows(self):
        return [r for r in _build_laporan_rows(self.db, False) if r['No_Kontrak'] == self.f.contract.no_kontrak]

    def payment(self, number='TEST-REPORT-P', amount=200, when=date(2026,10,1)):
        pay = models.Pembayaran(no_pembayaran=number, no_invoice=self.f.invoice.no_invoice,
            tanggal_pembayaran=when, nominal_transfer=amount, is_pph_disetor='false')
        self.db.add(pay); self.db.flush(); self.db.expire_all()
        return pay

    def test_cross_month_partial_ba(self):
        first = self.f.save_ba(volume=30)
        second = self.f.save_ba(no='TEST-REPORT-OCT', volume=20)
        second.tanggal_ba=date(2026,10,1); self.db.flush()
        actual = [r for r in self.rows() if r['Row_Type']=='REALISASI']
        self.assertEqual({r['Report_Date']:r['Volume_Pengambilan'] for r in actual}, {'2026-09-30':30,'2026-10-01':20})
        self.assertEqual(sum(r['Pendapatan_Pokok'] for r in actual),500)
        self.assertEqual(actual[0]['Outstanding_Pengambilan'],30)

    def test_payung_unbilled_ba_visible_no_fake_cash(self):
        self.f.contract.tipe_alur='PAYUNG_BA'; self.db.flush()
        from api.r_ba import create_ba
        ba=create_ba(schemas.BeritaAcaraCreate(no_ba='TEST-REPORT-PAYUNG-BA',no_kontrak=self.f.contract.no_kontrak,
            tanggal_ba=date(2026,9,30),volume_ba=30,harga_satuan=10,status='Selesai'),self.db,None)
        rows=[r for r in self.rows() if r['No_BA']==ba.no_ba]
        self.assertEqual(len(rows),1)
        self.assertEqual(rows[0]['Jumlah_Transfer'],0)
        self.assertEqual(rows[0]['Jumlah_Invoice'],0)
        self.assertEqual(rows[0]['Volume_Pengambilan'],30)
        self.assertEqual(rows[0]['Outstanding_Pengambilan'],0)

    def test_cash_in_includes_payment_without_do(self):
        first=self.payment('TEST-REPORT-P1',300,date(2026,9,1))
        self.payment('TEST-REPORT-P2',200)
        self.f.do.no_pembayaran=first.no_pembayaran; self.f.do.nominal_transfer=300
        self.db.flush();self.db.expire_all()
        rows=self.rows()
        self.assertEqual(sum(r['Jumlah_Transfer'] for r in rows),500)
        self.assertEqual(sum(r['Jumlah_Transfer'] for r in rows if r['Report_Date']=='2026-10-01'),200)
        self.assertEqual(sum(r['Pendapatan_Pokok'] for r in rows),800)

    def test_same_date_payment_attaches_once(self):
        self.f.save_ba(volume=30)
        self.payment('TEST-REPORT-P1',100,date(2026,9,30))
        self.payment('TEST-REPORT-P2',200,date(2026,9,30))
        rows=self.rows()
        self.assertEqual(sum(r['Jumlah_Transfer'] for r in rows),300)
        self.assertEqual(sum(r['Volume_Pengambilan'] or 0 for r in rows),30)
        self.assertEqual(sum(r['Pendapatan_Pokok'] for r in rows),300)

    def test_planned_date_not_transfer_no_ba(self):
        self.payment(when=date(2026,10,1))
        rows=self.rows()
        planned=next(r for r in rows if r['Row_Type']=='RENCANA')
        self.assertEqual(planned['Report_Date'],'2026-09-03')
        self.assertIsNone(planned['Volume_Pengambilan'])
        self.assertEqual(planned['Pendapatan_Pokok'],800)

    def test_ba_cannot_be_reused_by_second_do(self):
        ba=self.f.save_ba(volume=30)
        pay=self.payment()
        data=schemas.DeliveryOrderCreate(no_do='TEST-REPORT-DO2',no_pembayaran=pay.no_pembayaran,
            tanggal_do=date(2026,9,3),volume_do=20,no_ba=ba.no_ba)
        with self.assertRaises(HTTPException): create_do(data,self.db,None)

    def test_paid_invoice_cannot_be_deleted(self):
        self.payment()
        with patch('api.r_do.reverse_stok_do'):delete_do(self.f.do.no_do,self.db,None)
        with self.assertRaises(HTTPException):delete_invoice(self.f.invoice.no_invoice,self.db,None)
        self.assertEqual(self.db.query(models.Pembayaran).filter_by(no_pembayaran='TEST-REPORT-P').count(),1)

    def test_pph_tracking_allowed_after_superman_cash_still_locked(self):
        pay=self.payment()
        data=schemas.PembayaranCreate(no_pembayaran=pay.no_pembayaran,no_invoice=pay.no_invoice,
            tanggal_pembayaran=pay.tanggal_pembayaran,nominal_transfer=pay.nominal_transfer,is_pph_disetor='true')
        saved=create_pembayaran(data,self.db,None)
        self.assertEqual(saved['is_pph_disetor'],'true')
        with self.assertRaises(HTTPException):create_pembayaran(data.model_copy(update={'nominal_transfer':201}),self.db,None)

    def test_document_flow_normal_and_payung(self):
        ba=self.f.save_ba(volume=30)
        flow=build_document_flow(self.db,self.f.contract)
        self.assertIn({'source':'DO:'+self.f.do.no_do,'target':'BA:'+ba.no_ba},flow['edges'])
        self.assertEqual(next(n for n in flow['nodes'] if n['id']=='BA:'+ba.no_ba)['unit_price'],10)
        self.f.contract.tipe_alur='PAYUNG_BA'
        ba.no_do=None;self.f.invoice.no_ba=ba.no_ba;self.db.flush()
        flow=build_document_flow(self.db,self.f.contract)
        self.assertIn({'source':'KONTRAK:'+self.f.contract.no_kontrak,'target':'BA:'+ba.no_ba},flow['edges'])
        self.assertIn({'source':'BA:'+ba.no_ba,'target':'INVOICE:'+self.f.invoice.no_invoice},flow['edges'])
        self.assertIsNone(next(n for n in flow['nodes'] if n['kind']=='KONTRAK')['unit_price'])
        self.assertEqual(next(n for n in flow['nodes'] if n['kind']=='INVOICE')['unit_price'],10)

    def test_duplicate_records_flagged_not_silently_deleted(self):
        self.f.save_ba(volume=30)
        self.f.save_ba(no='TEST-REPORT-DUP',volume=30)
        rows=[r for r in self.rows() if r['Row_Type']=='REALISASI']
        self.assertEqual(len(rows),2)
        self.assertTrue(all(any('ganda' in w for w in r['Reporting_Warnings']) for r in rows))

    def test_ho_uses_same_sales_source(self):
        self.f.save_ba(volume=30)
        self.payment()
        content=generate_laporan_ho_xlsx([],year='2026',month='09',db=self.db,filters={'units':['TEST-LINK-UNIT']})
        wb=load_workbook(io.BytesIO(content));ws=wb.active
        # Total local value across fixed commodity rows equals canonical sales.
        value=sum(float(ws.cell(r,6).value or 0) for r in LOKAL_ROW_BY_KEY.values())
        self.assertEqual(value,300)

    def test_completed_ba_cannot_revert_to_draft(self):
        self.f.save_ba(volume=30)
        with self.assertRaises(HTTPException):self.f.save_ba(volume=30,status='Draft')


if __name__=='__main__':unittest.main()
